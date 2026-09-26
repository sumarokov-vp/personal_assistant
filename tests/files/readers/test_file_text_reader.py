import io
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from src.files.readers.file_text_reader import FileTextReader
from src.files.readers.unreadable_format_error import UnreadableFormatError

ITINERARY = (
    Path(__file__).parents[2]
    / "dropbox"
    / "fixtures"
    / "Itinerary_ALA_CNX_12-11-2026_000000000000.pdf"
)


def _docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Договор аренды")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Сумма"
    table.cell(0, 1).text = "1000"
    table.cell(1, 0).text = "Срок"
    table.cell(1, 1).text = "12 мес"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = "Расходы"
    first.append(["Дата", "Сумма"])
    first.append(["01.09.2026", 150])
    second = workbook.create_sheet("Доходы")
    second.append(["Зарплата", None, 5000])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _blank_pdf_bytes() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _read(content: bytes, name: str, reader: FileTextReader | None = None) -> str:
    return (reader or FileTextReader()).read(io.BytesIO(content), name).text


def test_docx_paragraphs_then_table_rows():
    assert _read(_docx_bytes(), "contract.DOCX") == (
        "Договор аренды\n\nСумма\t1000\nСрок\t12 мес"
    )


def test_xlsx_sheets_with_tab_separated_rows():
    assert _read(_xlsx_bytes(), "budget.xlsx") == (
        "## Расходы\nДата\tСумма\n01.09.2026\t150\n\n## Доходы\nЗарплата\t\t5000"
    )


def test_text_pdf_is_read():
    text = _read(ITINERARY.read_bytes(), ITINERARY.name)

    assert "12.11.2026" in text
    assert "CNX" in text


def test_pdf_without_text_layer_points_to_scan():
    with pytest.raises(UnreadableFormatError, match="скан"):
        _read(_blank_pdf_bytes(), "scan.pdf")


@pytest.mark.parametrize("name", ["old.xls", "old.doc", "photo.heic", "noext"])
def test_unsupported_format_lists_readable_ones(name: str):
    with pytest.raises(UnreadableFormatError, match="docx, xlsx"):
        _read(b"data", name)


def test_long_text_is_truncated():
    extracted = FileTextReader(max_text_chars=5).read(
        io.BytesIO(_docx_bytes()), "contract.docx"
    )

    assert extracted.text == "Догов"
    assert extracted.truncated


def test_oversized_document_is_refused():
    with pytest.raises(UnreadableFormatError, match="больше"):
        _read(_docx_bytes(), "contract.docx", FileTextReader(max_document_bytes=10))
