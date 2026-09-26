import io

from docx import Document


class DocxTextExtractor:
    def extract(self, content: bytes, name: str, max_chars: int) -> str:
        document = Document(io.BytesIO(content))
        lines = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]
        for table in document.tables:
            lines.append("")
            lines.extend(
                "\t".join(cell.text for cell in row.cells) for row in table.rows
            )
            if sum(len(line) for line in lines) > max_chars:
                break
        return "\n".join(lines)
