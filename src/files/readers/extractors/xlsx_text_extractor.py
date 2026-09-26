import io
from collections.abc import Iterable

from openpyxl import load_workbook


class XlsxTextExtractor:
    def extract(self, content: bytes, name: str, max_chars: int) -> str:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        lines: list[str] = []
        length = 0
        for sheet in workbook.worksheets:
            if lines:
                lines.append("")
            lines.append(f"## {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                line = _row_line(row)
                if not line.strip():
                    continue
                lines.append(line)
                length += len(line) + 1
                if length > max_chars:
                    break
            if length > max_chars:
                break
        workbook.close()
        return "\n".join(lines)


def _row_line(row: Iterable[object]) -> str:
    return "\t".join("" if value is None else str(value) for value in row)
